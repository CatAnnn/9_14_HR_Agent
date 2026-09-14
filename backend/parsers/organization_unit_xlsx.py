from __future__ import annotations

import hashlib
import html
import re
import unicodedata
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from backend.exceptions.parser_errors import ParserError


class OrganizationUnitXlsxParser:
    """Parse the organization workbook without flattening department rows."""

    REQUIRED_HEADERS = (
        "Organization",
        "Department Level 1",
        "Department Level 2",
        "Department Level 3",
        "Department Path",
        "Content",
    )
    CONTRACT_VERSION = "organization-unit-xlsx-v5-clean-content-groups"
    CHUNK_SIZE = 1600
    CHUNK_OVERLAP = 160

    _ORGANIZATION_ALIASES = {
        "Corporate Headquarters and Service Areas": [
            "公司总部与服务领域",
            "Corporate",
            "C/xx",
            "C xx",
        ],
        "Mobility (BBM)": ["移动出行", "Mobility", "BBM", "M/xx", "M xx"],
        "Industrial Technology (BBI)": ["工业技术", "Industrial Technology", "BBI"],
        "Consumer Goods (BBG)": ["消费品", "Consumer Goods", "BBG"],
        "Energy and Building Technology (BBE)": [
            "能源与建筑技术",
            "Energy and Building Technology",
            "BBE",
        ],
        "Value Accelerator & Portfolio Companies (VP)": [
            "价值加速器与投资组合公司",
            "Value Accelerator",
            "Portfolio Companies",
            "VP",
        ],
    }
    _TOP_LEVEL_CODE_ALIASES = {
        "BD": ["BD/xx", "BD xx", "Bosch Digital", "博世数字化"],
        "CR": ["CR/xx", "CR xx", "Corporate Research", "企业研究"],
        "GS": ["GS/xx", "GS xx", "Global Services", "全球服务"],
        "BBM": ["M/xx", "M xx", "Mobility", "移动出行"],
    }
    _NOISE_LINES = {
        "anonymized information",
        "back",
        "back to top",
        "bosch connect content is loading.",
        "bosch connect is not available",
        "close",
        "content provided by bosch connect",
        "the requested bosch connect community does not exists or does not provide this feature.",
        "user has no bosch connect profile.",
        "contact",
        "contact pageowner",
        "download",
        "home",
        "menu",
        "page owner",
        "print",
        "quick link to",
        "quick links",
        "search",
        "share",
        "skip to content",
        "to top",
    }
    _WIDGET_STRONG_LINES = {
        "bgn news service currently not available.",
        "bosch connect blogbeiträge werden geladen.",
        "bosch connect blogposts are loading.",
        "bosch connect content could not be loaded.",
        "bosch connect content is loading.",
        "bosch connect inhalt konnte nicht geladen werden.",
        "bosch connect inhalt wird geladen.",
        "bosch connect is not available",
        "content provided by bosch connect",
        "der user hat kein bosch connect profil.",
        "die angeforderte bosch connect community existiert nicht oder stellt diese funktion nicht bereit.",
        "du hast kein zugriff auf diese bosch connect community",
        "error occured",
        "people liked this",
        "the requested bosch connect community does not exists or does not provide this feature.",
        "user has no bosch connect profile.",
        "you haven't access on this bosch connect community",
    }
    _WIDGET_WEAK_LINE_RADIUS = {
        "comment": 2,
        "comments": 2,
        "more news": 2,
        "news": 2,
        "today": 4,
        "yesterday": 4,
    }
    _CONTACT_FIELD_TYPES = {
        "e-mail": "email",
        "fax": "phone",
        "fax number": "phone",
        "faxnummer": "phone",
        "mobile phone": "phone",
        "mobile phone number": "phone",
        "mobiltelefon": "phone",
        "mobiltelefonnummer": "phone",
        "phone number": "phone",
        "telephone": "phone",
        "telefon": "phone",
        "telefonnummer": "phone",
    }
    _CONTACT_IMAGE_TERMS = (
        "address",
        "anschrift",
        "bosch.connect",
        "contact person",
        "fax",
        "kontaktperson",
        "mobile",
        "mobiltelefon",
        "phone",
        "profile",
        "telefon",
    )
    _ATTACHMENT_TYPE_PATTERN = re.compile(
        r"^(?:7z|csv|doc|docx|gif|jpeg|jpg|pdf|png|ppt|pptx|rar|svg|tif|tiff|txt|xls|xlsx|zip)$",
        re.I,
    )
    _IMAGE_LINE_PATTERN = re.compile(
        r"^\[image:\s*(?P<label>[^\]]+)\](?:\s*bosch\s*connect)?$",
        re.I,
    )
    _GENERIC_IMAGE_LABELS = {
        ".",
        "image",
        "undefined",
    }
    _TEMPLATE_PLACEHOLDER_LINE = re.compile(
        r"^(?:"
        r"\{\{\s*(?:date|article-(?:headline|content))\s*\}\}|"
        r"\[image:\s*\{\{\s*article-img-alt\s*\}\}\s*\]"
        r"(?:\s*bosch\s*connect)?"
        r")$",
        re.I,
    )
    _HTML_ENTITY_PATTERN = re.compile(
        r"&(?:nbsp|amp|quot|apos|lt|gt|#\d+|#x[0-9a-f]+);",
        re.I,
    )
    _PUNCTUATION_ONLY_PATTERN = re.compile(r"^[\s|:;,.\-–—_·•]+$")
    _PHONE_VALUE_PATTERN = re.compile(r"^[+()\[\]0-9\s./\-–—xXextEXT]+$")
    _EMAIL_VALUE_PATTERN = re.compile(
        r"^(?:mailto:)?[^\s@]+@[^\s@]+\.[^\s@]+$",
        re.I,
    )
    _NO_CONTENT_PATTERNS = (
        re.compile(
            r"^(?:sorry[, ]*)?(?:this\s+(?:page|site)|page)\s+is\s+"
            r"(?:(?:(?:only\s+)?ava(?:il|li)able(?:\s+only)?\s+in|"
            r"only\s+in)\s+german|only\s+in\s+german\s+"
            r"ava(?:il|li)able)(?:\s+language)?[!. ]*$",
            re.I,
        ),
        re.compile(
            r"^(?:the\s+english\s+version\s+of\s+this\s+page\s+is\s+not\s+"
            r"available[. ]*please\s+refer\s+to\s+the\s+german\s+page|"
            r"currently\s+english\s+content\s+not\s+available[, ]+"
            r"please\s+refer\s+german\s+pages?)[!. ]*$",
            re.I,
        ),
        re.compile(r"^<?(?:page|side)?\s*under\s+construction>?[!. ]*$", re.I),
        re.compile(
            r"^(?:in\s+bearbeitung|seite\s+befindet\s+sich.*(?:aufbau|bearbeitung))[!. ]*$",
            re.I,
        ),
    )
    _NAVIGATION_CHROME_PATTERN = re.compile(
        r"^(?:topics?|themen|topics?\s+a-z|themen\s+a-z|"
        r"back\s+to\s+main\s+topics?.*|alphabetische\s+themenliste)$",
        re.I,
    )
    _DEPRECATED_PATTERN = re.compile(r"(?:^|[_\s-])to[_\s-]*be[_\s-]*deleted(?:$|[_\s-])", re.I)
    _SLASH_CODE_PATTERN = re.compile(
        r"(?<![A-Za-z0-9-])"
        r"(?P<code>[A-Z][A-Za-z0-9-]{0,31}"
        r"(?:/[A-Za-z0-9][A-Za-z0-9-]{0,31})+)"
        r"(?![A-Za-z0-9-])"
    )
    _PAREN_CODE_PATTERN = re.compile(r"\(([A-Z][A-Z0-9-]{1,15})\)")
    _SIMPLE_CODE_PATTERN = re.compile(r"^[A-Z][A-Z0-9-]{1,15}$")
    _WHITESPACE_PATTERN = re.compile(r"[ \t\f\v]+")

    def __init__(self) -> None:
        self.last_metadata: dict[str, Any] = {}

    @classmethod
    def matches(cls, path: Path) -> bool:
        if path.suffix.lower() != ".xlsx":
            return False
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception:  # noqa: BLE001
            return False
        try:
            return cls._matching_sheet(workbook) is not None
        finally:
            workbook.close()

    def parse(self, path: Path) -> str:
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise ParserError(f"组织架构 Excel 无法打开: {path}: {exc}") from exc

        try:
            sheet = self._matching_sheet(workbook)
            if sheet is None:
                raise ParserError(
                    f"组织架构 Excel 缺少固定表头 {list(self.REQUIRED_HEADERS)}: {path}"
                )
            records, stats = self._parse_sheet(sheet)
        finally:
            workbook.close()

        if not records:
            raise ParserError(f"组织架构 Excel 没有可索引记录: {path}")

        text = "\n\n".join(record["search_text"] for record in records)
        self.last_metadata = {
            "input_path": str(path),
            "sheet": sheet.title,
            "record_type": "organization_unit",
            "parser_contract": self.CONTRACT_VERSION,
            "structured_records": records,
            "structured_record_count": len(records),
            "chunk_size": self.CHUNK_SIZE,
            "chunk_overlap": self.CHUNK_OVERLAP,
            **stats,
        }
        return text

    @classmethod
    def _matching_sheet(cls, workbook: Any) -> Any | None:
        candidates = []
        if "Pages" in workbook.sheetnames:
            candidates.append(workbook["Pages"])
        candidates.extend(
            workbook[name] for name in workbook.sheetnames if name != "Pages"
        )
        required = set(cls.REQUIRED_HEADERS)
        for sheet in candidates:
            first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
            headers = {cls._cell(value) for value in first_row if cls._cell(value)}
            if required <= headers:
                return sheet
        return None

    def _parse_sheet(self, sheet: Any) -> tuple[list[dict[str, Any]], dict[str, int]]:
        first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        header_indexes = {
            self._cell(value): index
            for index, value in enumerate(first_row)
            if self._cell(value)
        }
        grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
        source_row_count = 0
        deprecated_row_count = 0

        for source_row, values in enumerate(
            sheet.iter_rows(min_row=2, values_only=True), start=2
        ):
            row = {
                header: self._cell(values[index] if index < len(values) else None)
                for header, index in header_indexes.items()
            }
            if not any(row.get(header) for header in self.REQUIRED_HEADERS):
                continue
            source_row_count += 1
            hierarchy_values = [
                row.get("Organization", ""),
                row.get("Department Level 1", ""),
                row.get("Department Level 2", ""),
                row.get("Department Level 3", ""),
                row.get("Department Path", ""),
            ]
            if self._is_deprecated(hierarchy_values):
                deprecated_row_count += 1
                continue

            record = self._build_record(row=row, source_row=source_row)
            grouped[record["metadata"]["record_identity_normalized"]].append(record)

        selected: list[dict[str, Any]] = []
        duplicate_row_count = 0
        for candidates in grouped.values():
            duplicate_row_count += max(0, len(candidates) - 1)
            best = max(candidates, key=self._record_preference)
            best["metadata"]["duplicate_source_rows"] = sorted(
                int(candidate["metadata"]["source_row"]) for candidate in candidates
            )
            best["metadata"]["duplicate_row_count"] = len(candidates) - 1
            selected.append(best)

        selected.sort(key=lambda item: item["metadata"]["department_path_normalized"])
        canonical_body_groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for item in selected:
            body_hash = str(
                item["metadata"].get("canonical_body_hash") or ""
            )
            if body_hash:
                canonical_body_groups[body_hash].append(item)
        for body_hash, group in canonical_body_groups.items():
            group_size = len(group)
            shared_group_id = body_hash if group_size > 1 else ""
            for item in group:
                item["metadata"]["canonical_body_group_id"] = shared_group_id
                item["metadata"]["canonical_body_group_size"] = group_size
        for item in selected:
            item["metadata"].setdefault("canonical_body_group_id", "")
            item["metadata"].setdefault("canonical_body_group_size", 0)

        empty_content_count = sum(
            item["metadata"]["status"] == "hierarchy_only" for item in selected
        )
        paths: dict[str, int] = defaultdict(int)
        for item in selected:
            paths[str(item["metadata"]["department_path_normalized"])] += 1
        cleanup_keys = (
            "removed_noise_line_count",
            "removed_widget_line_count",
            "removed_placeholder_line_count",
            "removed_contact_value_line_count",
            "removed_title_line_count",
            "removed_navigation_line_count",
            "removed_duplicate_line_count",
            "placeholder_content_removed_count",
            "removed_control_character_count",
            "decoded_html_entity_count",
        )
        return selected, {
            "source_row_count": source_row_count,
            "deprecated_row_count": deprecated_row_count,
            "duplicate_row_count": duplicate_row_count,
            "empty_content_count": empty_content_count,
            "canonical_body_group_count": sum(
                len(group) > 1 for group in canonical_body_groups.values()
            ),
            "canonical_body_group_record_count": sum(
                len(group)
                for group in canonical_body_groups.values()
                if len(group) > 1
            ),
            "path_collision_group_count": sum(count > 1 for count in paths.values()),
            "path_collision_record_count": sum(
                count for count in paths.values() if count > 1
            ),
            "excel_content_limit_count": sum(
                bool(item["metadata"].get("content_may_be_excel_truncated"))
                for item in selected
            ),
            **{
                key: sum(int(item["metadata"].get(key) or 0) for item in selected)
                for key in cleanup_keys
            },
        }

    def _build_record(self, *, row: dict[str, str], source_row: int) -> dict[str, Any]:
        organization = row.get("Organization", "")
        levels = [
            row.get("Department Level 1", ""),
            row.get("Department Level 2", ""),
            row.get("Department Level 3", ""),
        ]
        department_path = row.get("Department Path", "") or " > ".join(
            value for value in [organization, *levels] if value
        )
        path_segments = [
            self._cell(segment)
            for segment in re.split(r"\s*>\s*", department_path)
            if self._cell(segment)
        ]
        if not path_segments:
            path_segments = [value for value in [organization, *levels] if value]
        department_path = " > ".join(path_segments)
        canonical_path = self._canonical(department_path)

        source_title = row.get("Title", "")
        title = source_title or (
            path_segments[-1]
            if path_segments
            else organization or "Organization Unit"
        )
        page_topic = row.get("Page Topic", "")
        page_owner_department = row.get("Page Owner Department", "")
        raw_content = row.get("Content", "")
        content, cleanup_stats = self._clean_content_with_stats(
            raw_content,
            heading_values=[
                source_title,
                page_topic,
                path_segments[-1] if path_segments else "",
            ],
        )
        canonical_content = self._canonical(content)
        content_identity = hashlib.sha256(
            canonical_content.encode("utf-8")
        ).hexdigest()
        canonical_body_hash = content_identity if canonical_content else ""
        language = self._normalize_language(row.get("Language", ""), content)
        codes = self._extract_codes(
            [
                organization,
                *levels,
                department_path,
                page_owner_department,
            ]
        )
        organization_aliases = self._organization_aliases(organization)
        code_aliases = self._code_aliases(codes)
        context_prefix = self._context_prefix(
            title=title,
            page_topic=page_topic,
            department_path=department_path,
        )
        retrieval_prefix = self._retrieval_prefix(
            page_topic=page_topic,
            page_owner_department=page_owner_department,
            department_path=department_path,
            organization_aliases=organization_aliases,
            codes=codes,
        )

        status = "active" if content else "hierarchy_only"
        body = content or (
            "正文状态 / Content Status: 源文件未提供正文，"
            "或正文仅含不可检索占位内容；仅保留组织层级信息。"
        )
        source_url = row.get("URL", "")
        canonical_source_url = self._canonical(source_url)
        if canonical_source_url:
            identity_parts = [canonical_path, canonical_source_url]
        else:
            identity_parts = [
                canonical_path,
                self._canonical(title),
                self._canonical(page_topic),
                language,
                content_identity,
            ]
        record_identity = "\x1f".join(
            value for value in identity_parts if value
        )
        record_id = hashlib.sha256(record_identity.encode("utf-8")).hexdigest()
        canonical_path_id = hashlib.sha256(
            canonical_path.encode("utf-8")
        ).hexdigest()
        return {
            "record_id": record_id,
            "title": title,
            "context_prefix": context_prefix,
            "retrieval_prefix": retrieval_prefix,
            "content": content,
            "search_text": f"{context_prefix}\n\n{body}",
            "metadata": {
                "record_id": record_id,
                "canonical_path_id": canonical_path_id,
                "canonical_body_hash": canonical_body_hash,
                "record_identity_normalized": record_identity,
                "record_type": "organization_unit",
                "organization": organization,
                "organization_aliases": organization_aliases,
                "discovered_in_organizations": row.get(
                    "Discovered In Organizations",
                    "",
                ),
                "ownership_confidence": row.get("Ownership Confidence", ""),
                "department_level_1": levels[0],
                "department_level_2": levels[1],
                "department_level_3": levels[2],
                "department_path": department_path,
                "department_path_normalized": canonical_path,
                "path_segments": path_segments,
                "hierarchy_depth": len(path_segments),
                "department_codes": codes,
                "department_code_aliases": code_aliases,
                "page_topic": page_topic,
                "page_title": title,
                "page_owner_department": page_owner_department,
                "last_changed": row.get("Last Changed", ""),
                "source_url": source_url,
                "sidebar_complete": row.get("Sidebar Complete", ""),
                "source_row": source_row,
                "language": language,
                "status": status,
                "source_content_length": len(raw_content),
                "cleaned_content_length": len(content),
                "content_may_be_excel_truncated": len(raw_content) >= 32767,
                "parser_contract": self.CONTRACT_VERSION,
                **cleanup_stats,
            },
        }

    @classmethod
    def _context_prefix(
        cls,
        *,
        title: str,
        page_topic: str,
        department_path: str,
    ) -> str:
        rows = [f"完整组织路径 / Department Path: {department_path}"]
        if title:
            rows.append(f"页面标题 / Page Title: {title}")
        if page_topic and cls._canonical(page_topic) != cls._canonical(title):
            rows.append(f"页面主题 / Page Topic: {page_topic}")
        return "\n".join(rows)

    @classmethod
    def _retrieval_prefix(
        cls,
        *,
        page_topic: str,
        page_owner_department: str,
        department_path: str,
        organization_aliases: list[str],
        codes: list[str],
    ) -> str:
        rows: list[str] = []
        if page_topic:
            rows.append(f"Page topic: {page_topic}")
        if (
            page_owner_department
            and cls._canonical(page_owner_department)
            not in cls._canonical(department_path)
        ):
            rows.append(f"Page owner department: {page_owner_department}")
        aliases = cls._deduplicate(
            [
                *organization_aliases,
                *cls._code_aliases(codes[-2:]),
            ]
        )
        if aliases:
            rows.append("Search aliases: " + " | ".join(aliases))
        return "\n".join(rows)

    @classmethod
    def _extract_codes(cls, values: list[str]) -> list[str]:
        codes: list[str] = []
        for value in values:
            codes.extend(
                match.group("code")
                for match in cls._SLASH_CODE_PATTERN.finditer(value)
                if cls._looks_like_organization_code(match.group("code"))
            )
            codes.extend(cls._PAREN_CODE_PATTERN.findall(value))
            stripped = value.strip()
            if cls._SIMPLE_CODE_PATTERN.fullmatch(stripped):
                codes.append(stripped)
        normalized_codes = (
            cls._normalize_code(code)
            for code in codes
        )
        return cls._deduplicate(
            code for code in normalized_codes if code
        )

    @staticmethod
    def _normalize_code(value: str) -> str:
        normalized = unicodedata.normalize("NFKC", str(value or ""))
        return normalized.strip(" -/").upper()

    @staticmethod
    def _looks_like_organization_code(value: str) -> bool:
        segments = str(value or "").split("/")
        if len(segments) < 2 or any(
            not segment
            or segment.startswith("-")
            or segment.endswith("-")
            for segment in segments
        ):
            return False
        bases = [segment.split("-", 1)[0] for segment in segments]
        return (
            any(character.isdigit() for character in value)
            or any(
                len(base) >= 2 and base.isupper()
                for base in bases
            )
        )

    @classmethod
    def _code_aliases(cls, codes: list[str]) -> list[str]:
        aliases: list[str] = []
        for code in codes:
            aliases.extend(
                [
                    code.replace("/", " ").replace("-", " "),
                    code.replace("/", "").replace("-", ""),
                ]
            )
            aliases.extend(cls._TOP_LEVEL_CODE_ALIASES.get(code, []))
        return cls._deduplicate(alias for alias in aliases if alias and alias not in codes)

    @classmethod
    def _organization_aliases(cls, organization: str) -> list[str]:
        return cls._deduplicate(cls._ORGANIZATION_ALIASES.get(organization, []))

    @classmethod
    def _clean_content(cls, value: str) -> str:
        content, _ = cls._clean_content_with_stats(value)
        return content

    @classmethod
    def _clean_content_with_stats(
        cls,
        value: str,
        *,
        heading_values: list[str] | None = None,
    ) -> tuple[str, dict[str, int]]:
        source_text = str(value or "")
        decoded_html_entity_count = len(
            cls._HTML_ENTITY_PATTERN.findall(source_text)
        )
        normalized = unicodedata.normalize("NFKC", html.unescape(source_text))
        normalized_characters: list[str] = []
        removed_control_character_count = 0
        for character in normalized:
            if character in {"\r", "\n", "\t"}:
                normalized_characters.append(character)
                continue
            category = unicodedata.category(character)
            if category == "Cf":
                removed_control_character_count += 1
                continue
            if category == "Cc" or character == "\ufffd":
                removed_control_character_count += 1
                normalized_characters.append(" ")
                continue
            normalized_characters.append(character)
        normalized = "".join(normalized_characters)
        raw_lines = [
            cls._WHITESPACE_PATTERN.sub(" ", raw_line).strip()
            for raw_line in normalized.replace("\r\n", "\n")
            .replace("\r", "\n")
            .split("\n")
        ]
        alphabet_tokens = {
            line.casefold()
            for line in raw_lines
            if re.fullmatch(r"[a-zäöü]", line, re.I)
        }
        has_alphabetical_navigation = len(alphabet_tokens) >= 18
        heading_keys = {
            cls._canonical(item)
            for item in (heading_values or [])
            if cls._canonical(item)
        }
        widget_anchor_indexes = {
            index
            for index, line in enumerate(raw_lines)
            if (
                cls._canonical(line) in cls._WIDGET_STRONG_LINES
                or cls._TEMPLATE_PLACEHOLDER_LINE.fullmatch(line)
            )
        }
        stats = {
            "removed_noise_line_count": 0,
            "removed_widget_line_count": 0,
            "removed_placeholder_line_count": 0,
            "removed_contact_value_line_count": 0,
            "removed_title_line_count": 0,
            "removed_navigation_line_count": 0,
            "removed_duplicate_line_count": 0,
            "placeholder_content_removed_count": 0,
            "removed_control_character_count": (
                removed_control_character_count
            ),
            "decoded_html_entity_count": decoded_html_entity_count,
        }

        output: list[str] = []
        previous: str | None = None
        pending_contact_value: str | None = None
        meaningful_index = 0
        for source_index, line in enumerate(raw_lines):
            if not line:
                if output and output[-1] != "":
                    output.append("")
                continue
            canonical_line = cls._canonical(line)
            folded = canonical_line.rstrip(":").strip()

            if cls._TEMPLATE_PLACEHOLDER_LINE.fullmatch(line):
                stats["removed_noise_line_count"] += 1
                stats["removed_widget_line_count"] += 1
                stats["removed_placeholder_line_count"] += 1
                continue
            if canonical_line in cls._WIDGET_STRONG_LINES:
                stats["removed_noise_line_count"] += 1
                stats["removed_widget_line_count"] += 1
                continue
            weak_radius = cls._WIDGET_WEAK_LINE_RADIUS.get(canonical_line)
            if weak_radius is not None and any(
                abs(source_index - anchor_index) <= weak_radius
                for anchor_index in widget_anchor_indexes
            ):
                stats["removed_noise_line_count"] += 1
                stats["removed_widget_line_count"] += 1
                continue
            if cls._is_no_content_text(line):
                stats["removed_noise_line_count"] += 1
                stats["removed_placeholder_line_count"] += 1
                continue

            if pending_contact_value:
                should_remove = (
                    pending_contact_value == "phone"
                    and cls._looks_like_phone_value(line)
                ) or (
                    pending_contact_value == "email"
                    and cls._looks_like_email_value(line)
                )
                pending_contact_value = None
                if should_remove:
                    stats["removed_contact_value_line_count"] += 1
                    continue

            image_match = cls._IMAGE_LINE_PATTERN.fullmatch(line)
            if image_match:
                image_label = cls._canonical(image_match.group("label"))
                if (
                    image_label in cls._GENERIC_IMAGE_LABELS
                    or any(
                        term in image_label
                        for term in cls._CONTACT_IMAGE_TERMS
                    )
                ):
                    if any(
                        term in image_label
                        for term in (
                            "phone",
                            "telefon",
                            "mobile",
                            "mobiltelefon",
                            "fax",
                        )
                    ):
                        pending_contact_value = "phone"
                    stats["removed_noise_line_count"] += 1
                    continue

            contact_type = cls._CONTACT_FIELD_TYPES.get(folded)
            if contact_type:
                pending_contact_value = contact_type
                stats["removed_noise_line_count"] += 1
                continue
            if (
                folded in cls._NOISE_LINES
                or cls._ATTACHMENT_TYPE_PATTERN.fullmatch(line)
            ):
                stats["removed_noise_line_count"] += 1
                continue
            if cls._PUNCTUATION_ONLY_PATTERN.fullmatch(line):
                stats["removed_noise_line_count"] += 1
                continue
            if has_alphabetical_navigation and (
                re.fullmatch(
                    r"(?:[a-zäöü]|0-9|all|nach oben|to top)",
                    line,
                    re.I,
                )
                or cls._NAVIGATION_CHROME_PATTERN.fullmatch(line)
            ):
                stats["removed_navigation_line_count"] += 1
                continue
            meaningful_index += 1
            if (
                meaningful_index <= 6
                and cls._canonical(line) in heading_keys
            ):
                stats["removed_title_line_count"] += 1
                continue
            if line == previous:
                stats["removed_duplicate_line_count"] += 1
                continue
            output.append(line)
            previous = line

        content = "\n".join(output).strip()
        flattened = " ".join(content.split())
        if flattened and cls._is_no_content_text(flattened):
            content = ""
            stats["placeholder_content_removed_count"] = 1
        elif not content and stats["removed_placeholder_line_count"]:
            stats["placeholder_content_removed_count"] = 1
        return content, stats

    @classmethod
    def _is_no_content_text(cls, value: str) -> bool:
        flattened = " ".join(str(value or "").split())
        return bool(flattened) and any(
            pattern.fullmatch(flattened)
            for pattern in cls._NO_CONTENT_PATTERNS
        )

    @classmethod
    def _looks_like_phone_value(cls, value: str) -> bool:
        return (
            len(re.findall(r"\d", value)) >= 5
            and bool(cls._PHONE_VALUE_PATTERN.fullmatch(value.strip()))
        )

    @classmethod
    def _looks_like_email_value(cls, value: str) -> bool:
        return bool(cls._EMAIL_VALUE_PATTERN.fullmatch(value.strip()))

    @classmethod
    def _normalize_language(cls, value: str, content: str) -> str:
        normalized = cls._canonical(value)
        if normalized in {"de", "de-de", "deutsch", "german"}:
            return "de"
        if normalized in {"en", "en-us", "en-gb", "english"}:
            return "en"
        if normalized in {"zh", "zh-cn", "chinese", "中文"}:
            return "zh"
        return cls._detect_language(content)

    @classmethod
    def _detect_language(cls, content: str) -> str:
        if not content:
            return "unknown"
        if re.search(r"[\u3400-\u9fff]", content):
            return "zh"
        words = re.findall(r"[a-zA-ZÀ-ÿ]+", content.casefold())
        german = sum(word in {"der", "die", "das", "für", "mit", "und", "von", "wir"} for word in words)
        english = sum(word in {"and", "for", "of", "the", "to", "we", "with"} for word in words)
        if german > english:
            return "de"
        if english:
            return "en"
        return "unknown"

    @classmethod
    def _record_preference(
        cls,
        record: dict[str, Any],
    ) -> tuple[int, int, int, str]:
        language_rank = {"en": 4, "zh": 3, "unknown": 2, "de": 1}
        content = str(record.get("content") or "")
        return (
            int(bool(content)),
            language_rank.get(str(record["metadata"].get("language")), 0),
            len(content),
            content,
        )

    @classmethod
    def _is_deprecated(cls, values: list[str]) -> bool:
        return any(cls._DEPRECATED_PATTERN.search(value or "") for value in values)

    @classmethod
    def _cell(cls, value: Any) -> str:
        if value is None:
            return ""
        return cls._WHITESPACE_PATTERN.sub(" ", str(value)).strip()

    @staticmethod
    def _canonical(value: str) -> str:
        normalized = unicodedata.normalize("NFKC", value).casefold()
        return " ".join(normalized.split())

    @staticmethod
    def _deduplicate(values: Any) -> list[str]:
        output: list[str] = []
        seen: set[str] = set()
        for raw_value in values:
            value = str(raw_value or "").strip()
            key = value.casefold()
            if not value or key in seen:
                continue
            seen.add(key)
            output.append(value)
        return output
