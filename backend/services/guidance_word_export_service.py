from __future__ import annotations

import re
from dataclasses import dataclass
from io import BytesIO
from urllib.parse import quote

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

from backend.schemas.guidance import GuidanceReport
from backend.schemas.locale import SessionLocale, localized_value
from backend.schemas.retrieval import Citation, CitationAnchor


GUIDANCE_WORD_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
_CHINESE_FONT = "Microsoft YaHei"
_JAPANESE_FONT = "Yu Gothic"
_LATIN_FONT = "Arial"
_DIMENSIONS = (
    ("start", "开场定调与绩效结果对齐"),
    ("emotion", "情绪承接、接纳与共情"),
    ("requirement", "产出与标准"),
    ("plan", "总结与差异化发展计划"),
)
_DIMENSIONS_EN = (
    ("start", "Opening Tone and Performance-Outcome Alignment"),
    ("emotion", "Emotional Acknowledgement, Acceptance, and Empathy"),
    ("requirement", "Outcomes and Standards"),
    ("plan", "Summary and Differentiated Development Plan"),
)
_DIMENSIONS_DE = (
    ("start", "Gesprächseinstieg und Abstimmung des Leistungsergebnisses"),
    ("emotion", "Emotionale Aufnahme, Akzeptanz und Empathie"),
    ("requirement", "Ergebnisse und Standards"),
    ("plan", "Zusammenfassung und differenzierter Entwicklungsplan"),
)
_DIMENSIONS_JA = (
    ("start", "導入と評価結果のすり合わせ"),
    ("emotion", "感情の受け止め、受容と共感"),
    ("requirement", "成果と基準"),
    ("plan", "まとめと個別化した育成計画"),
)
_DIMENSIONS_BY_LOCALE = {
    "zh-CN": _DIMENSIONS,
    "en": _DIMENSIONS_EN,
    "de": _DIMENSIONS_DE,
    "ja": _DIMENSIONS_JA,
}
_CITATION_SCOPE_LABELS = {
    "career": "Career Elements",
    "culture": "文化与价值观",
    "development_dialog": "发展对话",
    "emotion": "情绪与沟通",
    "employee": "员工与岗位",
    "feedback": "反馈方法",
    "general": "通用管理",
    "job_level": "岗位与职级",
    "organization_unit": "组织信息",
    "performance": "绩效管理",
    "redline": "制度红线",
}
_CITATION_SCOPE_LABELS_EN = {
    "career": "Career Elements",
    "culture": "Culture and Values",
    "development_dialog": "Development Conversation",
    "emotion": "Emotion and Communication",
    "employee": "Employee and Role",
    "feedback": "Feedback Methods",
    "general": "General Management",
    "job_level": "Role and Job Level",
    "organization_unit": "Organization Information",
    "performance": "Performance Management",
    "redline": "Policy Red Lines",
}
_CITATION_SCOPE_LABELS_DE = {
    "career": "Career Elements",
    "culture": "Kultur und Werte",
    "development_dialog": "Entwicklungsgespräch",
    "emotion": "Emotion und Kommunikation",
    "employee": "Beschäftigte und Rolle",
    "feedback": "Feedbackmethoden",
    "general": "Allgemeines Management",
    "job_level": "Rolle und Stellenstufe",
    "organization_unit": "Organisationsinformationen",
    "performance": "Leistungsmanagement",
    "redline": "Verbindliche Richtlinien",
}
_CITATION_SCOPE_LABELS_JA = {
    "career": "Career Elements",
    "culture": "文化と価値観",
    "development_dialog": "育成対話",
    "emotion": "感情とコミュニケーション",
    "employee": "従業員と役割",
    "feedback": "フィードバック手法",
    "general": "一般的なマネジメント",
    "job_level": "役割と職位",
    "organization_unit": "組織情報",
    "performance": "パフォーマンス管理",
    "redline": "制度上の必須事項",
}
_CITATION_SCOPE_LABELS_BY_LOCALE = {
    "zh-CN": _CITATION_SCOPE_LABELS,
    "en": _CITATION_SCOPE_LABELS_EN,
    "de": _CITATION_SCOPE_LABELS_DE,
    "ja": _CITATION_SCOPE_LABELS_JA,
}
_LEGACY_SECTION_TITLES = {
    "zh-CN": ("沟通目的", "开场建议", "风险预判", "回应策略", "建议表达"),
    "en": (
        "Communication Purpose",
        "Opening Suggestion",
        "Risk Preview",
        "Response Strategies",
        "Suggested Phrasing",
    ),
    "de": (
        "Gesprächsziel",
        "Vorschlag für den Einstieg",
        "Risikoeinschätzung",
        "Reaktionsstrategien",
        "Formulierungsvorschläge",
    ),
    "ja": ("面談目的", "導入の提案", "リスクの見立て", "応答方針", "推奨表現"),
}
_WORD_REFERENCE_COPY = {
    "zh-CN": {
        "heading": "知识引用",
        "description": "以下内容为正文知识高亮对应的直接依据；仅有旧式引用时标记为相关原文。",
        "knowledge": "知识库",
        "fallback_scope": "知识资料",
        "direct": "直接依据：",
        "related": "相关原文：",
    },
    "en": {
        "heading": "Knowledge References",
        "description": (
            "The following source excerpts support highlighted passages in the "
            "document; legacy-only references are marked as related source text."
        ),
        "knowledge": "Knowledge Base",
        "fallback_scope": "Knowledge Material",
        "direct": "Direct evidence: ",
        "related": "Related source text: ",
    },
    "de": {
        "heading": "Wissensquellen",
        "description": (
            "Die folgenden Quellenauszüge belegen hervorgehobene Textstellen; reine "
            "Altreferenzen sind als zugehöriger Quelltext gekennzeichnet."
        ),
        "knowledge": "Wissensbasis",
        "fallback_scope": "Wissensmaterial",
        "direct": "Direkter Beleg: ",
        "related": "Zugehöriger Quelltext: ",
    },
    "ja": {
        "heading": "知識ソース",
        "description": (
            "以下の出典抜粋は本文で強調された箇所の直接的な根拠です。旧形式のみの引用は"
            "関連原文として表示します。"
        ),
        "knowledge": "知識ベース",
        "fallback_scope": "知識資料",
        "direct": "直接の根拠：",
        "related": "関連原文：",
    },
}
_WORD_DOCUMENT_COPY = {
    "zh-CN": ("谈前指导", "面谈前沟通准备建议", "谈前指导"),
    "en": (
        "Pre-conversation Guidance",
        "Preparation Guidance for the Performance Conversation",
        "Pre-conversation_Guidance",
    ),
    "de": (
        "Gesprächsvorbereitung",
        "Vorbereitungshinweise für den Leistungsdialog",
        "Gespraechsvorbereitung",
    ),
    "ja": ("面談前ガイダンス", "人事評価面談の準備ガイダンス", "面談前ガイダンス"),
}
_WORD_FONTS = {
    "zh-CN": _CHINESE_FONT,
    "en": _LATIN_FONT,
    "de": _LATIN_FONT,
    "ja": _JAPANESE_FONT,
}
_MAX_WORD_SOURCE_QUOTE_CHARACTERS = 1200


@dataclass(frozen=True)
class _TargetCitation:
    number: int
    citation: Citation
    anchor: CitationAnchor | None = None


@dataclass
class _GuidanceCitationIndex:
    sources: list[tuple[int, Citation]]
    by_target: dict[str, list[_TargetCitation]]


def _citation_identity(citation: Citation) -> str:
    source_type = citation.source_type or "knowledge_base"
    source_identity = citation.chunk_id or ":".join(
        (citation.source_id, citation.title, citation.scope)
    )
    return f"{source_type}:{source_identity}"


def _build_citation_index(report: GuidanceReport) -> _GuidanceCitationIndex:
    source_numbers: dict[str, int] = {}
    sources: list[tuple[int, Citation]] = []
    by_target: dict[str, list[_TargetCitation]] = {}
    for citation in report.citations:
        if not _source_quotes(citation):
            continue
        identity = _citation_identity(citation)
        number = source_numbers.get(identity)
        if number is None:
            number = len(sources) + 1
            source_numbers[identity] = number
            sources.append((number, citation))

        anchor_targets: set[str] = set()
        for anchor in citation.anchors:
            anchor_targets.add(anchor.target)
            by_target.setdefault(anchor.target, []).append(
                _TargetCitation(number=number, citation=citation, anchor=anchor)
            )
        for target in citation.targets:
            if target in anchor_targets:
                continue
            by_target.setdefault(target, []).append(
                _TargetCitation(number=number, citation=citation)
            )
    return _GuidanceCitationIndex(sources=sources, by_target=by_target)


def _add_reference_numbers(paragraph: Paragraph, numbers: list[int]) -> None:
    if not numbers:
        return
    run = paragraph.add_run(f" [{', '.join(str(number) for number in numbers)}]")
    run.font.superscript = True
    run.font.color.rgb = RGBColor(0, 110, 173)


def _add_cited_paragraph(
    document: DocxDocument,
    text: str,
    *,
    target: str,
    citation_index: _GuidanceCitationIndex,
    style: str | None = None,
) -> Paragraph:
    paragraph = document.add_paragraph(style=style)
    references = citation_index.by_target.get(target, [])
    if not references:
        paragraph.add_run(text)
        return paragraph

    grouped_matches: dict[tuple[int, int], set[int]] = {}
    unmatched_numbers: set[int] = set()
    for reference in references:
        highlight_text = reference.anchor.highlight_text if reference.anchor else ""
        start = text.find(highlight_text) if highlight_text else -1
        if start < 0:
            unmatched_numbers.add(reference.number)
            continue
        grouped_matches.setdefault(
            (start, start + len(highlight_text)), set()
        ).add(reference.number)

    selected_matches: list[tuple[int, int, list[int]]] = []
    cursor = 0
    for (start, end), numbers in sorted(
        grouped_matches.items(),
        key=lambda item: (item[0][0], -(item[0][1] - item[0][0])),
    ):
        if start < cursor:
            unmatched_numbers.update(numbers)
            continue
        selected_matches.append((start, end, sorted(numbers)))
        cursor = end

    if not selected_matches:
        paragraph.add_run(text)
        _add_reference_numbers(paragraph, sorted(unmatched_numbers))
        return paragraph

    cursor = 0
    cited_numbers: set[int] = set()
    for start, end, numbers in selected_matches:
        if start > cursor:
            paragraph.add_run(text[cursor:start])
        highlighted = paragraph.add_run(text[start:end])
        highlighted.font.highlight_color = WD_COLOR_INDEX.YELLOW
        _add_reference_numbers(paragraph, numbers)
        cited_numbers.update(numbers)
        cursor = end
    if cursor < len(text):
        paragraph.add_run(text[cursor:])

    trailing_numbers = sorted(unmatched_numbers - cited_numbers)
    _add_reference_numbers(paragraph, trailing_numbers)
    return paragraph


def _set_style_font(
    style,
    *,
    size: float,
    bold: bool = False,
    font_name: str = _CHINESE_FONT,
) -> None:
    style.font.name = font_name
    style.font.size = Pt(size)
    style.font.bold = bold
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font_name)


def _configure_document(
    document: DocxDocument,
    *,
    locale: SessionLocale = "zh-CN",
) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    font_name = localized_value(locale, _WORD_FONTS)
    _set_style_font(document.styles["Normal"], size=10.5, font_name=font_name)
    _set_style_font(document.styles["Title"], size=22, bold=True, font_name=font_name)
    _set_style_font(document.styles["Heading 1"], size=15, bold=True, font_name=font_name)
    _set_style_font(document.styles["Heading 2"], size=11.5, bold=True, font_name=font_name)
    _set_style_font(document.styles["List Bullet"], size=10.5, font_name=font_name)

    document.styles["Normal"].paragraph_format.space_after = Pt(5)
    document.styles["Normal"].paragraph_format.line_spacing = 1.35
    document.styles["Heading 1"].paragraph_format.space_before = Pt(14)
    document.styles["Heading 1"].paragraph_format.space_after = Pt(7)
    document.styles["Heading 2"].paragraph_format.space_before = Pt(8)
    document.styles["Heading 2"].paragraph_format.space_after = Pt(4)


def _add_dimension_points(
    document: DocxDocument,
    report: GuidanceReport,
    citation_index: _GuidanceCitationIndex,
) -> None:
    assert report.dimension_points is not None
    dimensions = localized_value(report.locale, _DIMENSIONS_BY_LOCALE)
    for field_name, title in dimensions:
        document.add_heading(title, level=1)
        for group_index, group in enumerate(
            getattr(report.dimension_points, field_name)
        ):
            target_prefix = f"dimension_points.{field_name}.{group_index}"
            document.add_heading(group.title, level=2)
            if group.summary:
                _add_cited_paragraph(
                    document,
                    group.summary,
                    target=f"{target_prefix}.summary",
                    citation_index=citation_index,
                )
            for detail_index, detail in enumerate(group.details):
                _add_cited_paragraph(
                    document,
                    detail,
                    target=f"{target_prefix}.details.{detail_index}",
                    citation_index=citation_index,
                    style="List Bullet",
                )


def _add_legacy_sections(
    document: DocxDocument,
    report: GuidanceReport,
    citation_index: _GuidanceCitationIndex,
) -> None:
    section_titles = localized_value(report.locale, _LEGACY_SECTION_TITLES)
    section_fields = (
        ("purpose", [report.purpose]),
        ("opening_suggestion", [report.opening_suggestion]),
        ("risk_preview", report.risk_preview),
        ("response_strategies", report.response_strategies),
        ("safer_phrases", report.safer_phrases),
    )
    for title, (target_prefix, items) in zip(
        section_titles,
        section_fields,
        strict=True,
    ):
        document.add_heading(title, level=1)
        for item_index, item in enumerate(items):
            target = (
                target_prefix
                if target_prefix in {"purpose", "opening_suggestion"}
                else f"{target_prefix}.{item_index}"
            )
            _add_cited_paragraph(
                document,
                item,
                target=target,
                citation_index=citation_index,
                style="List Bullet",
            )


def _source_quotes(citation: Citation) -> list[tuple[str, bool]]:
    quotes: list[tuple[str, bool]] = []
    seen: set[str] = set()
    for anchor in citation.anchors:
        source_quote = anchor.source_quote.strip()
        if source_quote and source_quote not in seen:
            seen.add(source_quote)
            quotes.append((source_quote, True))
    if not quotes and citation.quote:
        related_quote = citation.quote.strip()
        if related_quote:
            quotes.append((related_quote, False))
    return quotes


def _limited_source_quote(value: str) -> str:
    if len(value) <= _MAX_WORD_SOURCE_QUOTE_CHARACTERS:
        return value
    return value[: _MAX_WORD_SOURCE_QUOTE_CHARACTERS - 1].rstrip() + "…"


def _add_citation_appendix(
    document: DocxDocument,
    citation_index: _GuidanceCitationIndex,
    *,
    locale: SessionLocale = "zh-CN",
) -> None:
    sources = [
        (number, citation, _source_quotes(citation))
        for number, citation in citation_index.sources
    ]
    sources = [entry for entry in sources if entry[2]]
    if not sources:
        return

    document.add_page_break()
    copy = localized_value(locale, _WORD_REFERENCE_COPY)
    document.add_heading(copy["heading"], level=1)
    document.add_paragraph(copy["description"])
    scope_labels = localized_value(locale, _CITATION_SCOPE_LABELS_BY_LOCALE)
    knowledge_label = copy["knowledge"]
    fallback_scope = copy["fallback_scope"]
    for number, citation, quotes in sources:
        heading = document.add_paragraph()
        heading.paragraph_format.space_before = Pt(8)
        heading.add_run(f"[{number}] {citation.title or citation.source_id}").bold = True
        scope = citation.scope.strip().lower().replace("-", "_")
        source_type = "Skill" if citation.source_type == "skill" else knowledge_label
        meta = document.add_paragraph(
            f"{source_type} · {scope_labels.get(scope, scope or fallback_scope)}"
        )
        meta.runs[0].font.size = Pt(9)
        meta.runs[0].font.color.rgb = RGBColor(91, 104, 113)
        for source_quote, exact in quotes:
            quote_paragraph = document.add_paragraph(style="List Bullet")
            label_text = copy["direct"] if exact else copy["related"]
            label = quote_paragraph.add_run(label_text)
            label.bold = True
            quote_run = quote_paragraph.add_run(
                _limited_source_quote(source_quote)
            )
            if exact:
                quote_run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def build_guidance_word_document(report: GuidanceReport) -> bytes:
    document = Document()
    _configure_document(document, locale=report.locale)
    citation_index = _build_citation_index(report)
    document_title, document_subject, _ = localized_value(
        report.locale,
        _WORD_DOCUMENT_COPY,
    )
    document.core_properties.title = document_title
    document.core_properties.subject = document_subject

    title = document.add_heading(document_title, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph(document_subject)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.color.rgb = RGBColor(91, 104, 113)
    subtitle.paragraph_format.space_after = Pt(14)

    if report.dimension_points is not None:
        _add_dimension_points(document, report, citation_index)
    else:
        _add_legacy_sections(document, report, citation_index)

    if report.disclaimer:
        disclaimer = document.add_paragraph()
        disclaimer.paragraph_format.space_before = Pt(14)
        run = disclaimer.add_run(report.disclaimer)
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(91, 104, 113)

    _add_citation_appendix(document, citation_index, locale=report.locale)

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def guidance_word_content_disposition(
    session_id: str,
    locale: SessionLocale = "zh-CN",
) -> str:
    safe_session_id = re.sub(r"[^A-Za-z0-9_-]+", "_", session_id).strip("_")
    safe_session_id = safe_session_id[:64] or "session"
    ascii_name = f"guidance_{safe_session_id}.docx"
    _, _, display_prefix = localized_value(locale, _WORD_DOCUMENT_COPY)
    display_name = f"{display_prefix}_{safe_session_id}.docx"
    return (
        f'attachment; filename="{ascii_name}"; '
        f"filename*=UTF-8''{quote(display_name)}"
    )
