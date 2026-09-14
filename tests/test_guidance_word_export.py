from io import BytesIO

import pytest
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from fastapi import FastAPI
from fastapi.testclient import TestClient

from backend.api.dependencies import (
    get_coach_service,
    get_guidance_service,
    get_session_service,
)
from backend.api.error_handlers import register_error_handlers
from backend.api.routes import guidance, reports
from backend.schemas.guidance import GuidanceReport
from backend.schemas.retrieval import Citation
from backend.services.guidance_word_export_service import (
    GUIDANCE_WORD_MEDIA_TYPE,
    build_guidance_word_document,
    guidance_word_content_disposition,
)


def _report() -> GuidanceReport:
    def groups(title: str, first: str, second: str):
        return [
            {
                "title": title,
                "summary": f"{title}的核心准备方向。",
                "details": [first, second],
            }
        ]

    return GuidanceReport(
        session_id="session-中文-1",
        intent_id="development",
        purpose="围绕本次沟通目标对齐事实和后续安排。",
        opening_suggestion="先说明沟通目的，再邀请员工补充自己的视角。",
        risk_preview=["员工可能对事实依据提出疑问。"],
        response_strategies=["先确认员工的观点，再逐项核对事实。"],
        safer_phrases=["我们先把双方掌握的信息放在一起核对。"],
        dimension_points={
            "start": groups(
                "明确沟通目标",
                "先说明本次沟通要解决的核心问题。",
                "明确事实核对与行动讨论的先后顺序。",
            ),
            "emotion": groups(
                "承接员工反应",
                "先复述员工表达的主要担忧。",
                "确认理解后再继续讨论事实和方案。",
            ),
            "requirement": groups(
                "聚焦核心诉求",
                "区分员工提出的事实问题与资源诉求。",
                "逐项确认哪些事项可以当场回应。",
            ),
            "plan": groups(
                "形成行动计划",
                "明确下一步行动、责任人与时间节点。",
                "约定后续检查进展和调整计划的方式。",
            ),
        },
        citations=[
            {
                "chunk_id": "plan-source:0",
                "source_id": "plan-source",
                "title": "发展计划方法",
                "scope": "development_dialog",
                "targets": ["dimension_points.plan.0.details.0"],
                "anchors": [
                    {
                        "target": "dimension_points.plan.0.details.0",
                        "highlight_text": "明确下一步行动、责任人与时间节点",
                        "source_quote": "行动计划应明确责任人与时间节点。",
                        "source_context": "共同制定行动计划时，应明确责任人与时间节点，并约定复盘方式。",
                    }
                ],
            }
        ],
        disclaimer="本建议用于演练准备，不替代 HR/Legal 或 Manager 的最终判断。",
    )


def _document_text(payload: bytes) -> str:
    document = Document(BytesIO(payload))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def test_build_guidance_word_document_contains_visible_guidance() -> None:
    payload = build_guidance_word_document(_report())

    assert payload.startswith(b"PK")
    text = _document_text(payload)
    assert "谈前指导" in text
    assert "开场定调与绩效结果对齐" in text
    assert "明确沟通目标" in text
    assert "明确沟通目标的核心准备方向。" in text
    assert "情绪承接、接纳与共情" in text
    assert "产出与标准" in text
    assert "总结与差异化发展计划" in text
    assert "形成行动计划" in text
    assert "不替代 HR/Legal" in text
    assert "development" not in text


def test_english_guidance_word_document_localizes_only_its_shell() -> None:
    report = _report().model_copy(
        update={
            "locale": "en",
            "disclaimer": "Guidance for rehearsal only.",
        }
    )

    payload = build_guidance_word_document(report)
    document = Document(BytesIO(payload))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert document.core_properties.title == "Pre-conversation Guidance"
    assert "Opening Tone and Performance-Outcome Alignment" in text
    assert "Knowledge References" in text
    assert "Knowledge Base · Development Conversation" in text
    assert "Direct evidence: 行动计划应明确责任人与时间节点。" in text
    assert "谈前指导" not in text
    assert "Pre-conversation_Guidance_session-1.docx" in guidance_word_content_disposition(
        "session-1",
        "en",
    )


@pytest.mark.parametrize(
    ("locale", "title", "dimension", "references", "filename"),
    [
        (
            "de",
            "Gesprächsvorbereitung",
            "Gesprächseinstieg und Abstimmung des Leistungsergebnisses",
            "Wissensquellen",
            "Gespraechsvorbereitung_session-1.docx",
        ),
        (
            "ja",
            "面談前ガイダンス",
            "導入と評価結果のすり合わせ",
            "知識ソース",
            "%E9%9D%A2%E8%AB%87%E5%89%8D%E3%82%AC%E3%82%A4%E3%83%80%E3%83%B3%E3%82%B9_session-1.docx",
        ),
    ],
)
def test_guidance_word_document_localizes_german_and_japanese_shell(
    locale: str,
    title: str,
    dimension: str,
    references: str,
    filename: str,
) -> None:
    report = _report().model_copy(update={"locale": locale})

    payload = build_guidance_word_document(report)
    document = Document(BytesIO(payload))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert document.core_properties.title == title
    assert dimension in text
    assert references in text
    assert filename in guidance_word_content_disposition("session-1", locale)


def test_build_guidance_word_document_highlights_citations_and_adds_sources() -> None:
    document = Document(BytesIO(build_guidance_word_document(_report())))

    cited_paragraph = next(
        paragraph
        for paragraph in document.paragraphs
        if "明确下一步行动、责任人与时间节点" in paragraph.text
    )
    assert "[1]" in cited_paragraph.text
    assert any(
        "明确下一步行动、责任人与时间节点" in run.text
        and run.font.highlight_color == WD_COLOR_INDEX.YELLOW
        for run in cited_paragraph.runs
    )
    assert any(run.font.superscript for run in cited_paragraph.runs if "[1]" in run.text)

    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "知识引用" in text
    assert "[1] 发展计划方法" in text
    source_paragraph = next(
        paragraph
        for paragraph in document.paragraphs
        if "行动计划应明确责任人与时间节点" in paragraph.text
    )
    assert any(
        "行动计划应明确责任人与时间节点" in run.text
        and run.font.highlight_color == WD_COLOR_INDEX.YELLOW
        for run in source_paragraph.runs
    )


def test_guidance_word_skips_target_only_source_without_appendix_text() -> None:
    report = _report()
    report.citations.insert(
        0,
        Citation(
            chunk_id="missing-quote:0",
            source_id="missing-quote",
            title="缺少原文的旧引用",
            scope="general",
            targets=["dimension_points.plan.0.details.1"],
        ),
    )

    document = Document(BytesIO(build_guidance_word_document(report)))
    uncited_paragraph = next(
        paragraph
        for paragraph in document.paragraphs
        if "约定后续检查进展和调整计划的方式" in paragraph.text
    )
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert "[" not in uncited_paragraph.text
    assert "缺少原文的旧引用" not in text
    assert "[1] 发展计划方法" in text


def test_guidance_word_download_endpoint_returns_docx() -> None:
    report = _report()

    class GuidanceServiceStub:
        @staticmethod
        def get(session_id: str) -> GuidanceReport:
            assert session_id == report.session_id
            return report

    class SessionServiceStub:
        @staticmethod
        def get_session(session_id: str):
            assert session_id == report.session_id
            return object()

    app = FastAPI()
    app.include_router(guidance.router)
    app.dependency_overrides[get_guidance_service] = GuidanceServiceStub
    app.dependency_overrides[get_session_service] = SessionServiceStub

    with TestClient(app) as client:
        response = client.get(f"/guidance/{report.session_id}/export.docx")

    assert response.status_code == 200
    assert response.headers["content-type"] == GUIDANCE_WORD_MEDIA_TYPE
    assert "filename*=UTF-8''" in response.headers["content-disposition"]
    assert response.headers["cache-control"] == "no-store"
    assert "承接员工反应" in _document_text(response.content)


def test_cached_report_reads_require_owned_session_before_repository_access() -> None:
    class OwnerFilteredSessionService:
        def __init__(self) -> None:
            self.calls: list[str] = []

        def get_session(self, session_id: str):
            self.calls.append(session_id)
            raise KeyError(f"Session not found: {session_id}")

    class CachedReportServiceSpy:
        def __init__(self) -> None:
            self.calls: list[str] = []

        def get(self, session_id: str):
            self.calls.append(session_id)
            raise AssertionError("cached report must not be read for an unowned session")

    for path in (
        "/guidance/other-users-session",
        "/guidance/other-users-session/export.docx",
        "/reports/other-users-session/coach",
    ):
        session_service = OwnerFilteredSessionService()
        report_service = CachedReportServiceSpy()
        app = FastAPI()
        register_error_handlers(app)
        app.include_router(guidance.router)
        app.include_router(reports.router)
        app.dependency_overrides[get_session_service] = lambda: session_service
        app.dependency_overrides[get_guidance_service] = lambda: report_service
        app.dependency_overrides[get_coach_service] = lambda: report_service

        with TestClient(app) as client:
            response = client.get(path)

        assert response.status_code == 404
        assert session_service.calls == ["other-users-session"]
        assert report_service.calls == []
