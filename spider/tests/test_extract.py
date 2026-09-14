from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import extract
import project_config


def page_record(
    organization: str,
    url: str,
    **overrides: object,
) -> dict[str, object]:
    record: dict[str, object] = {
        "url": url,
        "requested_url": url,
        "title": "Shared Services",
        "root_department": organization,
        "department_path": [organization, "Shared Services"],
        "language": "en",
        "content": "Contact support for shared services.",
        "sidebar_complete": True,
    }
    record.update(overrides)
    return record


def wcms_content(
    body: str,
    *,
    last_changed: str = "Jul 28, 2026",
    owner: str = "Jane Doe, M/HR",
    german: bool = False,
    trailing_lines: tuple[str, ...] = (),
) -> str:
    if german:
        toolbar = (
            "Redaktionswerkzeuge",
            "Zugriffsrechte sind erforderlich",
            "Prep Areas einblenden",
            "Prep Areas ausblenden",
            "Report generieren",
            "WCMS Hilfe",
            "SiteArchitect",
            "Seite bearbeiten",
            "Empfehlen / Kurz-ID",
            "Kurz-ID kopieren",
            "https://bgn.bosch.com/FIRSTspiritWeb/permlink/example-DE",
            "Seite empfehlen",
            owner,
        )
    else:
        toolbar = (
            "Editorial tools",
            "Explicit permissions necessary",
            "Show Prep Areas",
            "Hide Prep Areas",
            "Generate report",
            "WCMS help",
            "SiteArchitect",
            "Edit page",
            "localized recommendation label",
            "Copy short ID",
            "https://bgn.bosch.com/FIRSTspiritWeb/permlink/example-EN",
            "Recommend page",
            owner,
        )
    return "\n".join((body, last_changed, *toolbar, *trailing_lines))


def write_jsonl(records: list[dict[str, object]]) -> Path:
    temporary = tempfile.NamedTemporaryFile(
        mode="w",
        encoding="utf-8",
        suffix=".jsonl",
        delete=False,
    )
    with temporary:
        for record in records:
            temporary.write(json.dumps(record) + "\n")
    return Path(temporary.name)


class ExtractConfigurationTests(unittest.TestCase):
    def test_shared_project_configuration_is_used(self) -> None:
        self.assertEqual(
            tuple(extract.TARGET_ORGANIZATIONS),
            project_config.TARGET_ORGANIZATIONS,
        )
        self.assertEqual(extract.INPUT_PAGES_FILE, project_config.PAGES_FILE)
        self.assertEqual(
            extract.OUTPUT_EXCEL_FILE,
            project_config.EXCEL_OUTPUT_FILE,
        )

    def test_output_headers_are_the_compact_export_contract(self) -> None:
        expected = [
            "Organization",
            "Discovered In Organizations",
            "Ownership Confidence",
            "Department Level 1",
            "Department Level 2",
            "Department Level 3",
            "Page Topic",
            "Title",
            "Department Path",
            "Content",
        ]
        self.assertEqual(extract.OUTPUT_HEADERS, expected)
        for removed_header in (
            "Page Owner Department",
            "Last Changed",
            "Language",
            "URL",
            "Sidebar Complete",
        ):
            self.assertNotIn(removed_header, extract.OUTPUT_HEADERS)
        extract.validate_config()

    def test_zero_content_scan_limit_uses_full_content(self) -> None:
        with patch.object(extract, "TOPIC_CONTENT_SCAN_LENGTH", 0):
            category = extract.topic_category(
                "Unrelated topic",
                "Unrelated title",
                ["Unrelated department"],
                "Please contact support for assistance.",
            )
        self.assertEqual(category, "contacts and support")

    def test_excel_formulas_are_escaped(self) -> None:
        self.assertEqual(extract.excel_value("=1+1"), "'=1+1")

    def test_missing_enabled_filter_field_fails_fast(self) -> None:
        path = write_jsonl(
            [page_record(project_config.TARGET_ORGANIZATIONS[0], "https://bgn.bosch.com/a")]
        )
        self.addCleanup(path.unlink, missing_ok=True)
        with patch.object(extract, "ALLOWED_HTTP_STATUS", [200]):
            with self.assertRaisesRegex(ValueError, "http_status"):
                extract.read_data(path)


class ContentCleanupTests(unittest.TestCase):
    def _read_one(self, record: dict[str, object]) -> dict[str, object]:
        path = write_jsonl([record])
        self.addCleanup(path.unlink, missing_ok=True)
        result = extract.read_data(path)
        self.assertEqual(len(result.rows), 1)
        return result.rows[0]

    def test_english_toolbar_is_removed_and_metadata_is_recovered(self) -> None:
        row = self._read_one(
            page_record(
                project_config.TARGET_ORGANIZATIONS[0],
                "https://bgn.bosch.com/body.html",
                content=wcms_content("Useful body"),
            )
        )
        self.assertEqual(row["Content"], "Useful body")
        self.assertEqual(row["_last_changed"], "Jul 28, 2026")
        self.assertEqual(row["_page_owner_department"], "M/HR")

    def test_german_toolbar_and_trailing_site_footer_are_removed(self) -> None:
        row = self._read_one(
            page_record(
                project_config.TARGET_ORGANIZATIONS[0],
                "https://bgn.bosch.com/de.html",
                language="de",
                content=wcms_content(
                    "Nützlicher Inhalt",
                    last_changed="28.07.2026",
                    owner="Max Mustermann, C/HR",
                    german=True,
                    trailing_lines=("About Bosch GlobalNet", "Compliance"),
                ),
            )
        )
        self.assertEqual(row["Content"], "Nützlicher Inhalt")
        self.assertEqual(row["_last_changed"], "28.07.2026")
        self.assertEqual(row["_page_owner_department"], "C/HR")

    def test_explicit_metadata_wins_over_footer_fallback(self) -> None:
        row = self._read_one(
            page_record(
                project_config.TARGET_ORGANIZATIONS[0],
                "https://bgn.bosch.com/explicit.html",
                content=wcms_content("Useful body"),
                last_changed="2026-08-01",
                page_owner_department="C/Explicit",
            )
        )
        self.assertEqual(row["_last_changed"], "2026-08-01")
        self.assertEqual(row["_page_owner_department"], "C/Explicit")

    def test_incomplete_toolbar_signature_is_preserved(self) -> None:
        content = "Body\nEditorial tools\nWCMS help\nNot a real footer"
        metadata = extract.parse_content_metadata(content)
        self.assertFalse(metadata.toolbar_removed)
        self.assertEqual(metadata.content, content)

    def test_empty_body_after_footer_cleanup_is_excluded(self) -> None:
        path = write_jsonl(
            [
                page_record(
                    project_config.TARGET_ORGANIZATIONS[0],
                    "https://bgn.bosch.com/empty.html",
                    content=wcms_content(""),
                )
            ]
        )
        self.addCleanup(path.unlink, missing_ok=True)
        result = extract.read_data(path)
        self.assertEqual(result.rows, [])
        self.assertEqual(result.excluded, 1)

    def test_toolbar_help_text_does_not_change_topic_category(self) -> None:
        row = self._read_one(
            page_record(
                project_config.TARGET_ORGANIZATIONS[0],
                "https://bgn.bosch.com/unrelated.html",
                title="Unrelated title",
                department_path=[
                    project_config.TARGET_ORGANIZATIONS[0],
                    "Unrelated department",
                ],
                content=wcms_content("Neutral information"),
            )
        )
        self.assertTrue(str(row["Page Topic"]).endswith("overview and key information"))


class ExtractDeduplicationTests(unittest.TestCase):
    def test_same_page_is_globally_deduplicated_and_sources_are_retained(self) -> None:
        url = "https://bgn.bosch.com/shared.html"
        organizations = project_config.TARGET_ORGANIZATIONS[:2]
        path = write_jsonl([page_record(org, url) for org in organizations])
        self.addCleanup(path.unlink, missing_ok=True)

        result = extract.read_data(path)

        self.assertEqual(len(result.rows), 1)
        self.assertEqual(result.duplicates, 1)
        self.assertEqual(
            result.rows[0]["Discovered In Organizations"],
            " | ".join(organizations),
        )
        self.assertEqual(result.rows[0]["Ownership Confidence"], "Low (ambiguous)")

    def test_url_owner_beats_a_shorter_unrelated_path(self) -> None:
        corporate, mobility = project_config.TARGET_ORGANIZATIONS[:2]
        url = (
            "https://bgn.bosch.com/FIRSTspiritWeb/wcms/wcms_c/en/"
            "bosch-globalnet/organization/mobility-bbm/topic.html"
        )
        records = [
            page_record(corporate, url, department_path=[corporate, "Topic"]),
            page_record(
                mobility,
                url,
                department_path=[mobility, "Division", "Department", "Topic"],
            ),
        ]
        path = write_jsonl(records)
        self.addCleanup(path.unlink, missing_ok=True)

        row = extract.read_data(path).rows[0]

        self.assertEqual(row["Organization"], mobility)
        self.assertEqual(row["Ownership Confidence"], "High (URL match)")

    def test_same_url_is_deduplicated_within_one_organization(self) -> None:
        url = "https://bgn.bosch.com/duplicate.html"
        record = page_record(project_config.TARGET_ORGANIZATIONS[0], url)
        path = write_jsonl([record, dict(record)])
        self.addCleanup(path.unlink, missing_ok=True)

        result = extract.read_data(path)

        self.assertEqual(len(result.rows), 1)
        self.assertEqual(result.duplicates, 1)

    def test_recommendation_alias_is_merged_with_plain_url(self) -> None:
        url = "https://bgn.bosch.com/page.html"
        alias = url + "?cmcall=true&perm_link=wcms_c_-page-en"
        organization = project_config.TARGET_ORGANIZATIONS[0]
        path = write_jsonl(
            [
                page_record(organization, alias, content="alias"),
                page_record(organization, url, content="plain"),
            ]
        )
        self.addCleanup(path.unlink, missing_ok=True)

        result = extract.read_data(path)

        self.assertEqual(len(result.rows), 1)
        self.assertEqual(result.rows[0]["_input_url"], url)
        self.assertEqual(result.rows[0]["Content"], "plain")

    def test_list_and_text_department_paths_share_the_same_identity(self) -> None:
        organization = project_config.TARGET_ORGANIZATIONS[0]
        records = [
            page_record(
                organization,
                "",
                requested_url="",
                title="Policy",
                department_path=[organization, " Division ", "", "Team"],
            ),
            page_record(
                organization,
                "",
                requested_url="",
                title="Policy",
                department_path=f"{organization} > Division > Team",
            ),
            page_record(
                organization,
                "",
                requested_url="",
                title="Policy",
                department_path=f"{organization} > Division > Other",
            ),
        ]
        path = write_jsonl(records)
        self.addCleanup(path.unlink, missing_ok=True)

        result = extract.read_data(path)

        self.assertEqual(result.duplicates, 1)
        self.assertEqual(len(result.rows), 2)
        rows = {row["Department Level 2"]: row for row in result.rows}
        self.assertEqual(set(rows), {"Other", "Team"})
        self.assertEqual(rows["Team"]["Department Level 1"], "Division")
        self.assertEqual(
            rows["Team"]["Department Path"],
            f"{organization} > Division > Team",
        )

    def test_business_query_is_not_removed(self) -> None:
        url = "https://bgn.bosch.com/page.html"
        organization = project_config.TARGET_ORGANIZATIONS[0]
        path = write_jsonl(
            [
                page_record(organization, url + "?view=one"),
                page_record(organization, url + "?view=two"),
            ]
        )
        self.addCleanup(path.unlink, missing_ok=True)
        self.assertEqual(len(extract.read_data(path).rows), 2)

    def test_same_path_in_different_languages_is_preserved(self) -> None:
        url = "https://bgn.bosch.com/page.html"
        organization = project_config.TARGET_ORGANIZATIONS[0]
        path = write_jsonl(
            [
                page_record(organization, url, language="en"),
                page_record(organization, url, language="de"),
            ]
        )
        self.addCleanup(path.unlink, missing_ok=True)
        self.assertEqual(len(extract.read_data(path).rows), 2)

    def test_url_path_case_is_preserved_in_identity(self) -> None:
        organization = project_config.TARGET_ORGANIZATIONS[0]
        path = write_jsonl(
            [
                page_record(organization, "https://bgn.bosch.com/Page.html"),
                page_record(organization, "https://bgn.bosch.com/page.html"),
            ]
        )
        self.addCleanup(path.unlink, missing_ok=True)
        self.assertEqual(len(extract.read_data(path).rows), 2)

    def test_deduplication_does_not_depend_on_visible_url_column(self) -> None:
        organization = project_config.TARGET_ORGANIZATIONS[0]
        url = "https://bgn.bosch.com/page.html"
        path = write_jsonl(
            [page_record(organization, url), page_record(organization, url)]
        )
        self.addCleanup(path.unlink, missing_ok=True)
        result = extract.read_data(path)
        self.assertEqual(len(result.rows), 1)

    def test_first_and_last_refer_to_input_order(self) -> None:
        organization = project_config.TARGET_ORGANIZATIONS[0]
        url = "https://bgn.bosch.com/page.html"
        records = [
            page_record(organization, url, title="Zulu", content="first"),
            page_record(organization, url, title="Alpha", content="last"),
        ]
        path = write_jsonl(records)
        self.addCleanup(path.unlink, missing_ok=True)
        with patch.object(extract, "DUPLICATE_KEEP", "first"):
            self.assertEqual(extract.read_data(path).rows[0]["Content"], "first")
        with patch.object(extract, "DUPLICATE_KEEP", "last"):
            self.assertEqual(extract.read_data(path).rows[0]["Content"], "last")


class ExcelWriterTests(unittest.TestCase):
    def test_write_is_atomic_and_uses_only_output_headers(self) -> None:
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "organization_unit.xlsx"
            row = {header: "" for header in extract.OUTPUT_HEADERS}
            row["_input_url"] = "https://bgn.bosch.com/page.html"
            row["_language"] = "en"
            row["_sidebar_complete"] = True
            row["_page_owner_department"] = "M/HR"
            row["_last_changed"] = "2026-08-01"
            row["Content"] = "Body"
            extract.write_excel(output, [row])

            workbook = load_workbook(output, read_only=False)
            try:
                sheet = workbook[extract.OUTPUT_SHEET_NAME]
                self.assertEqual(
                    [cell.value for cell in sheet[1]],
                    extract.OUTPUT_HEADERS,
                )
                self.assertEqual(sheet.max_column, len(extract.OUTPUT_HEADERS))
                self.assertFalse(sheet.sheet_view.showGridLines)
                self.assertEqual(
                    sheet.sheet_format.defaultRowHeight,
                    extract.BODY_ROW_HEIGHT,
                )
                self.assertNotIn(2, sheet.row_dimensions)
            finally:
                workbook.close()
            self.assertEqual(list(Path(directory).glob(".organization_unit.*.xlsx")), [])

    def test_failed_save_keeps_previous_file_and_removes_temporary(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "organization_unit.xlsx"
            output.write_bytes(b"previous workbook")
            with patch(
                "openpyxl.workbook.workbook.Workbook.save",
                side_effect=RuntimeError("save failed"),
            ):
                with self.assertRaisesRegex(RuntimeError, "save failed"):
                    extract.write_excel(output, [])
            self.assertEqual(output.read_bytes(), b"previous workbook")
            self.assertEqual(list(Path(directory).glob(".organization_unit.*.xlsx")), [])


class SummaryTests(unittest.TestCase):
    def test_incompatible_summary_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            pages = Path(directory) / "pages.jsonl"
            summary = Path(directory) / "summary.json"
            pages.write_text("", encoding="utf-8")
            summary.write_text(
                json.dumps({"output_schema_version": 4}),
                encoding="utf-8",
            )
            with self.assertRaises(ValueError):
                extract.load_crawl_summary(pages)


if __name__ == "__main__":
    unittest.main()
